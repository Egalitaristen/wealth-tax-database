#!/usr/bin/env python3
"""wtdb.py — helper for the Wealth Tax Database skill.

Standard library only, except `export`, which needs openpyxl (pip install openpyxl).

Commands
  init     DIR --area NAME          create the data folder and empty CSVs
  add      DIR TABLE JSON|@FILE     add rows (JSON object, JSON list, or @file with JSON lines); IDs assigned if blank
  set      DIR TABLE ID k=v ...     update fields on one row
  validate DIR                      check rules; exit code 1 if there are errors
  dedupe   DIR                      list likely duplicate actors
  merge    DIR KEEP_ID DROP_ID      merge a duplicate actor into another (moves evidence, links, profiles)
  checkquotes DIR [--run R] [--ids E1,E2] [--sample N] [--strict]
                                    fetch source pages and check each quote appears on the page
  status   DIR                      counts, coverage grid, sweep progress, next targets
  export   DIR [--out FILE]         write the xlsx report
"""
import argparse, csv, datetime as dt, difflib, json, os, re, sys, unicodedata
from collections import Counter, defaultdict

SCHEMA = {
    "actors": ["actor_id", "name", "kind", "actor_type", "other_types", "role", "affiliation", "party", "affiliation_actor_id", "country",
               "level", "location", "wikidata_qid", "website", "status", "alive_checked", "prominence", "notes",
               "added_run", "updated"],
    "profiles": ["profile_id", "actor_id", "platform", "url", "source", "checked"],
    "evidence": ["evidence_id", "actor_id", "date", "stance", "grade", "evidence_type", "wt_design", "design_detail",
                 "quote_original", "quote_lang", "quote_en", "source_url", "source_title", "publisher", "archive_url",
                 "accessed", "harvest_id", "grading_note", "verified", "run_id"],
    "geo_units": ["geo_unit_id", "name", "level", "unit_type", "parent_id", "website", "sweep_status", "swept_run",
                  "swept_date", "queries_run", "hits", "notes"],
    "links": ["link_id", "from_actor_id", "to_actor_id", "link_type", "evidence_id", "source_url", "note"],
    "harvest": ["harvest_id", "title", "list_type", "url", "date", "est_size", "status", "actors_added", "notes"],
    "exclusions": ["exclusion_id", "name", "reason", "detail", "source_url", "run_id"],
    "runs": ["run_id", "date", "assistant", "focus", "targets", "queries_run", "sources_opened", "candidates_seen",
             "new_actors", "new_evidence", "grade_upgrades", "verified_ok", "verified_failed", "saturation",
             "notes_next"],
}
PREFIX = {"actors": ("A", 4), "evidence": ("E", 5), "geo_units": ("G", 4), "links": ("L", 4),
          "harvest": ("H", 4), "exclusions": ("X", 4), "runs": ("R", 3), "profiles": ("P", 5)}
REQUIRED = {
    "actors": ["name", "kind", "actor_type", "country", "level", "status", "alive_checked"],
    "evidence": ["actor_id", "date", "stance", "grade", "evidence_type", "wt_design", "quote_original",
                 "quote_lang", "source_url", "grading_note"],
    "geo_units": ["name", "level", "sweep_status"],
    "links": ["from_actor_id", "to_actor_id", "link_type"],
    "profiles": ["actor_id", "platform", "url"],
    "harvest": ["title", "list_type", "status"],
    "exclusions": ["name", "reason"],
    "runs": ["date", "focus"],
}
ENUMS = {
    ("actors", "kind"): {"person", "organisation", "institution"},
    ("actors", "actor_type"): {"politician", "party", "government_body", "academic", "economist", "think_tank",
                               "ngo_campaign", "union", "union_official", "journalist", "media_outlet",
                               "legal_tax_professional", "wealthy_individual", "business_leader", "faith",
                               "cultural_figure", "campaigner", "international_body", "funder"},
    ("actors", "level"): {"supranational", "national", "regional", "local"},
    ("actors", "status"): {"active", "former_role", "dissolved"},
    ("actors", "alive_checked"): {"y", "na"},
    ("actors", "prominence"): {"", "international", "national", "regional", "local"},
    ("evidence", "stance"): {"support", "retracted"},
    ("evidence", "grade"): {"A", "B", "C"},
    ("evidence", "evidence_type"): {"vote", "bill_sponsorship", "motion", "motion_signature", "open_letter",
                                    "publication", "report", "op_ed", "editorial", "speech", "statement",
                                    "interview", "social_post", "manifesto", "policy_adoption", "campaign",
                                    "coalition_membership", "testimony", "funding", "other"},
    ("evidence", "wt_design"): {"annual_net_wealth", "one_off_levy", "billionaire_minimum", "unrealised_gains",
                                "exit_tax", "partial_wealth", "defend_existing", "wealth_tax_unspecified"},
    ("evidence", "verified"): {"", "n", "y", "failed"},
    ("geo_units", "level"): {"supranational", "national", "devolved", "region", "county", "combined_authority",
                             "council", "sub_council"},
    ("geo_units", "sweep_status"): {"pending", "in_progress", "done", "skipped"},
    ("links", "link_type"): {"member_of", "works_for", "leads", "co_author", "co_signatory", "funds", "founded",
                             "coalition_partner", "party_of"},
    ("harvest", "list_type"): {"roll_call", "cosponsors", "motion_signatories", "open_letter", "coalition_members",
                               "report_authors", "commission_evidence", "conference_motion", "other"},
    ("profiles", "platform"): {"website", "official_page", "x", "bluesky", "mastodon", "instagram", "facebook",
                               "youtube", "tiktok", "linkedin", "substack", "wikipedia"},
    ("harvest", "status"): {"queued", "in_progress", "done", "failed"},
    ("exclusions", "reason"): {"out_of_definition", "pre_2010_only", "deceased", "private_individual", "opposes",
                               "insufficient_evidence", "duplicate"},
}
PLATFORM_DOMAINS = {"x": ("x.com", "twitter.com"), "bluesky": ("bsky.app",), "instagram": ("instagram.com",),
                    "facebook": ("facebook.com", "fb.com"), "youtube": ("youtube.com", "youtu.be"),
                    "tiktok": ("tiktok.com",), "linkedin": ("linkedin.com",), "substack": ("substack.com",),
                    "wikipedia": ("wikipedia.org",)}
FORMAL_TYPES = {"vote", "bill_sponsorship", "motion", "motion_signature", "open_letter", "publication", "report",
                "manifesto", "policy_adoption", "testimony", "funding"}
GRADE_RANK = {"A": 3, "B": 2, "C": 1}
GRADE_LABEL = {"A": "explicit", "B": "named", "C": "appears to support"}
DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")
MIN_YEAR = 2010


# ---------- io ----------
def path(d, t):
    return os.path.join(d, t + ".csv")


def load(d, t):
    p = path(d, t)
    if not os.path.exists(p):
        return []
    with open(p, newline="", encoding="utf-8") as f:
        return [{k: (v or "").strip() for k, v in r.items() if k is not None} for r in csv.DictReader(f)]


def save(d, t, rows):
    p = path(d, t)
    tmp = p + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=SCHEMA[t], extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in SCHEMA[t]})
    os.replace(tmp, p)


def load_all(d):
    return {t: load(d, t) for t in SCHEMA}


def idcol(t):
    return SCHEMA[t][0]


def next_id(rows, t):
    pre, width = PREFIX[t]
    n = 0
    for r in rows:
        m = re.match(re.escape(pre) + r"(\d+)$", r.get(idcol(t), ""))
        if m:
            n = max(n, int(m.group(1)))
    return f"{pre}{n + 1:0{width}d}"


def norm_name(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(the|dr|prof|professor|sir|dame|lord|lady|baroness|cllr|councillor|mp|msp|ms|mr|mrs|rt hon)\b",
               " ", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def year(d):
    return int(d[:4]) if DATE_RE.match(d or "") else None


def need_dir(d):
    if not os.path.isdir(d):
        sys.exit(f"No data folder at {d}. Run: wtdb.py init {d} --area NAME")


# ---------- commands ----------
def cmd_init(a):
    os.makedirs(a.dir, exist_ok=True)
    for t in SCHEMA:
        if not os.path.exists(path(a.dir, t)):
            save(a.dir, t, [])
    prof = os.path.join(a.dir, "area_profile.md")
    if not os.path.exists(prof):
        tpl = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "area-profile-template.md")
        text = open(tpl, encoding="utf-8").read() if os.path.exists(tpl) else "# Area profile — {AREA}\n"
        with open(prof, "w", encoding="utf-8") as f:
            f.write(text.replace("{AREA}", a.area))
    meta = os.path.join(a.dir, "meta.json")
    if not os.path.exists(meta):
        slug = os.path.basename(os.path.normpath(a.dir))
        with open(meta, "w", encoding="utf-8") as f:
            json.dump({"name": a.area, "iso2": slug.split("-")[0].upper(), "lead": "", "status": "in progress",
                       "boundaries": "", "notes": ""}, f, indent=2, ensure_ascii=False)
            f.write("\n")
    print(f"Initialised {a.dir} for {a.area}")


def parse_rows(arg):
    if arg.startswith("@"):
        text = open(arg[1:], encoding="utf-8").read().strip()
        if text.startswith("["):
            return json.loads(text)
        return [json.loads(l) for l in text.splitlines() if l.strip()]
    v = json.loads(arg)
    return v if isinstance(v, list) else [v]


def cmd_add(a):
    need_dir(a.dir)
    t = a.table
    if t not in SCHEMA:
        sys.exit(f"Unknown table {t}. Tables: {', '.join(SCHEMA)}")
    rows = load(a.dir, t)
    existing = {r[idcol(t)] for r in rows}
    today = dt.date.today().isoformat()
    new_ids = []
    for r in parse_rows(a.rows):
        unknown = set(r) - set(SCHEMA[t])
        if unknown:
            sys.exit(f"Unknown column(s) for {t}: {', '.join(sorted(unknown))}")
        r = {k: ("" if v is None else str(v)) for k, v in r.items()}
        if not r.get(idcol(t)):
            r[idcol(t)] = next_id(rows, t)
        elif r[idcol(t)] in existing:
            sys.exit(f"ID {r[idcol(t)]} already exists in {t}")
        if t == "actors":
            r.setdefault("updated", today)
            r.setdefault("status", "active")
            if r.get("kind") in ("organisation", "institution"):
                r.setdefault("alive_checked", "na")
        if t == "evidence":
            r.setdefault("accessed", today)
            r.setdefault("verified", "n")
            r.setdefault("stance", "support")
        if t == "geo_units":
            r.setdefault("sweep_status", "pending")
        if t == "harvest":
            r.setdefault("status", "queued")
        if t == "runs":
            r.setdefault("date", today)
        rows.append(r)
        existing.add(r[idcol(t)])
        new_ids.append(r[idcol(t)])
    save(a.dir, t, rows)
    print("Added " + ", ".join(new_ids))


def cmd_set(a):
    need_dir(a.dir)
    rows = load(a.dir, a.table)
    hit = [r for r in rows if r.get(idcol(a.table)) == a.id]
    if not hit:
        sys.exit(f"{a.id} not found in {a.table}")
    for kv in a.pairs:
        if "=" not in kv:
            sys.exit(f"Expected key=value, got {kv}")
        k, v = kv.split("=", 1)
        if k not in SCHEMA[a.table]:
            sys.exit(f"Unknown column {k} for {a.table}")
        hit[0][k] = v
    if a.table == "actors":
        hit[0]["updated"] = dt.date.today().isoformat()
    save(a.dir, a.table, rows)
    print(f"Updated {a.id}")


def validate(db):
    errs, warns = [], []
    ids = {}
    for t, rows in db.items():
        seen = set()
        for i, r in enumerate(rows, start=2):
            rid = r.get(idcol(t), "")
            where = f"{t}.csv line {i} ({rid or 'no id'})"
            if not rid:
                errs.append(f"{where}: missing {idcol(t)}")
            elif rid in seen:
                errs.append(f"{where}: duplicate ID")
            seen.add(rid)
            for c in REQUIRED[t]:
                if not r.get(c):
                    errs.append(f"{where}: missing {c}")
            for (tt, c), allowed in ENUMS.items():
                if tt == t and r.get(c) and r[c] not in allowed:
                        errs.append(f"{where}: {c}='{r.get(c)}' not allowed ({', '.join(sorted(x for x in allowed if x))})")
        ids[t] = seen
    actors = {r["actor_id"]: r for r in db["actors"]}
    ev_by_actor = defaultdict(list)
    for i, e in enumerate(db["evidence"], start=2):
        where = f"evidence.csv line {i} ({e.get('evidence_id')})"
        ev_by_actor[e.get("actor_id")].append(e)
        if e.get("actor_id") and e["actor_id"] not in actors:
            errs.append(f"{where}: actor_id {e['actor_id']} does not exist")
        d = e.get("date", "")
        if d and not DATE_RE.match(d):
            errs.append(f"{where}: date '{d}' must be YYYY-MM-DD, YYYY-MM or YYYY")
        elif d and year(d) < MIN_YEAR:
            errs.append(f"{where}: date {d} is before {MIN_YEAR}")
        elif d and d > dt.date.today().isoformat():
            errs.append(f"{where}: date {d} is in the future")
        u = e.get("source_url", "")
        if u and not u.startswith(("http://", "https://")):
            errs.append(f"{where}: source_url must start with http(s)://")
        if e.get("quote_lang") and e["quote_lang"].lower() not in ("en", "eng") and not e.get("quote_en"):
            errs.append(f"{where}: quote_en (translation) needed for quote_lang={e['quote_lang']}")
        if e.get("harvest_id") and e["harvest_id"] not in ids["harvest"]:
            errs.append(f"{where}: harvest_id {e['harvest_id']} does not exist")
        if e.get("grade") == "A" and e.get("evidence_type") not in FORMAL_TYPES and \
                e.get("wt_design") == "wealth_tax_unspecified" and not e.get("design_detail"):
            warns.append(f"{where}: grade A on a {e.get('evidence_type')} with no design or detail — should this be B?")
        if e.get("grade") == "A" and e.get("evidence_type") not in FORMAL_TYPES and not e.get("design_detail") \
                and e.get("wt_design") != "wealth_tax_unspecified":
            warns.append(f"{where}: grade A {e.get('evidence_type')} — add design_detail (rate, threshold or bill)")
        if len(e.get("quote_original", "")) < 15:
            warns.append(f"{where}: very short quote")
        if e.get("verified") == "failed":
            errs.append(f"{where}: verification failed — fix or remove the row")
    geo = {g["geo_unit_id"] for g in db["geo_units"]}
    gidx = geo_index(db["geo_units"])
    names = defaultdict(list)
    qids = defaultdict(list)
    for i, r in enumerate(db["actors"], start=2):
        where = f"actors.csv line {i} ({r.get('actor_id')} {r.get('name')})"
        if r.get("actor_id") not in ev_by_actor:
            errs.append(f"{where}: no evidence — every actor needs at least one evidence row")
        if r.get("kind") == "person" and r.get("alive_checked") != "y":
            errs.append(f"{where}: person must have alive_checked=y")
        if r.get("kind") != "person" and r.get("alive_checked") == "y" and r.get("status") == "dissolved":
            warns.append(f"{where}: dissolved organisation marked alive_checked=y")
        uid, depth, rest = resolve_location(r.get("location", ""), gidx)
        if not r.get("location"):
            warns.append(f"{where}: no location (use at least the country name)")
        elif gidx[2] and not uid:
            errs.append(f"{where}: location must start with the country root unit "
                        f"({', '.join(gidx[0][u]['name'] for u in gidx[2].values())})")
        elif r.get("level") == "local" and depth <= 1 and not rest:
            warns.append(f"{where}: local actor located only at country level")
        if re.search(r"\balive\b|death record|no death|wikidata (id|match)|no wikidata", r.get("notes", ""), re.I):
            warns.append(f"{where}: notes narrate the alive or Wikidata check; alive_checked records it, so remove that from notes")
        for ot in type_list(r)[1:]:
            if ot not in ENUMS[("actors", "actor_type")]:
                errs.append(f"{where}: other_types has unknown type '{ot}'")
        if "politician" in type_list(r) and not r.get("party") and r.get("level") != "local":
            warns.append(f"{where}: politician without party (use the party's common name, or Independent; "
                         f"local officials may leave it blank when it is not known)")
        if r.get("affiliation_actor_id") and r["affiliation_actor_id"] not in actors:
            errs.append(f"{where}: affiliation_actor_id {r['affiliation_actor_id']} does not exist")
        if r.get("wikidata_qid") and not re.match(r"^Q\d+$", r["wikidata_qid"]):
            errs.append(f"{where}: wikidata_qid must look like Q12345")
        names[(norm_name(r.get("name", "")), r.get("kind"))].append(r.get("actor_id"))
        if r.get("wikidata_qid"):
            qids[r["wikidata_qid"]].append(r.get("actor_id"))
    for (n, k), v in names.items():
        if len(v) > 1:
            errs.append(f"actors: same name '{n}' ({k}) on {', '.join(v)} — merge, or add a distinguishing note "
                        f"and different names if they really are different people")
    for q, v in qids.items():
        if len(v) > 1:
            errs.append(f"actors: wikidata {q} used by {', '.join(v)} — duplicate")
    for i, l in enumerate(db["links"], start=2):
        for c in ("from_actor_id", "to_actor_id"):
            if l.get(c) and l[c] not in actors:
                errs.append(f"links.csv line {i}: {c} {l[c]} does not exist")
        if l.get("evidence_id") and l["evidence_id"] not in ids["evidence"]:
            errs.append(f"links.csv line {i}: evidence_id {l['evidence_id']} does not exist")
    seen_profiles = set()
    for i, p in enumerate(db["profiles"], start=2):
        where = f"profiles.csv line {i}"
        a = actors.get(p.get("actor_id"))
        if not a:
            errs.append(f"{where}: actor_id {p.get('actor_id')} does not exist")
        url = p.get("url", "")
        m = re.match(r"^https://([^/\s]+)(/\S*)?$", url)
        if not m:
            errs.append(f"{where}: url must be an https link (no mailto:, tel: or bare handles)")
        else:
            host = m.group(1).lower().split(":")[0]
            doms = PLATFORM_DOMAINS.get(p.get("platform"))
            if doms and not any(host == d or host.endswith("." + d) for d in doms):
                errs.append(f"{where}: {url} does not look like a {p.get('platform')} link")
        key = (p.get("actor_id"), url.rstrip("/").lower())
        if key in seen_profiles:
            warns.append(f"{where}: duplicate link for {p.get('actor_id')}")
        seen_profiles.add(key)
        if a and set(type_list(a)) == {"wealthy_individual"}:
            warns.append(f"{where}: {a.get('name')} is only a wealthy individual; link only official or professional pages of public figures")
    for i, g in enumerate(db["geo_units"], start=2):
        if g.get("parent_id") and g["parent_id"] not in geo:
            errs.append(f"geo_units.csv line {i}: parent_id {g['parent_id']} does not exist")
        if g.get("sweep_status") == "done" and not g.get("swept_date"):
            warns.append(f"geo_units.csv line {i} ({g.get('name')}): done without swept_date")
    return errs, warns


def cmd_validate(a):
    need_dir(a.dir)
    errs, warns = validate(load_all(a.dir))
    for w in warns:
        print("WARN  " + w)
    for e in errs:
        print("ERROR " + e)
    print(f"\n{len(errs)} error(s), {len(warns)} warning(s)")
    sys.exit(1 if errs else 0)


def cmd_dedupe(a):
    need_dir(a.dir)
    actors = load(a.dir, "actors")
    pairs = []
    for i in range(len(actors)):
        for j in range(i + 1, len(actors)):
            x, y = actors[i], actors[j]
            if x["kind"] != y["kind"]:
                continue
            nx, ny = norm_name(x["name"]), norm_name(y["name"])
            r = difflib.SequenceMatcher(None, nx, ny).ratio()
            px, py = nx.split(), ny.split()
            same_last = (x["kind"] == "person" and len(px) > 1 and len(py) > 1 and px[-1] == py[-1]
                         and px[0][0] == py[0][0])
            contains = x["kind"] != "person" and len(min(nx, ny, key=len)) > 4 and (nx in ny or ny in nx)
            if r >= 0.86 or same_last or contains:
                pairs.append((r, x, y))
    pairs.sort(key=lambda p: -p[0])
    for r, x, y in pairs:
        print(f"{r:.2f}  {x['actor_id']} {x['name']} ({x['actor_type']}, {x.get('affiliation','')})  <->  "
              f"{y['actor_id']} {y['name']} ({y['actor_type']}, {y.get('affiliation','')})")
    print(f"\n{len(pairs)} possible duplicate pair(s). Check each; merge true duplicates (keep the older ID).")


def derive(db):
    ev_by_actor = defaultdict(list)
    for e in db["evidence"]:
        ev_by_actor[e["actor_id"]].append(e)
    out = []
    for r in db["actors"]:
        evs = sorted(ev_by_actor.get(r["actor_id"], []), key=lambda e: e.get("date", ""))
        sup = [e for e in evs if e.get("stance") == "support"]
        latest = evs[-1] if evs else {}
        best = max((e.get("grade", "C") for e in sup), key=lambda g: GRADE_RANK.get(g, 0), default="")
        strongest = max(sup, key=lambda e: (GRADE_RANK.get(e.get("grade"), 0), e.get("date", "")), default={})
        out.append(dict(r,
                        best_grade=best,
                        support_label=GRADE_LABEL.get(best, ""),
                        support_status="retracted" if latest.get("stance") == "retracted" else "supporter",
                        evidence_count=str(len(sup)),
                        designs="; ".join(sorted({e.get("wt_design") for e in sup if e.get("wt_design")})),
                        evidence_types="; ".join(sorted({e.get("evidence_type") for e in sup})),
                        first_date=sup[0]["date"] if sup else "",
                        latest_date=sup[-1]["date"] if sup else "",
                        strongest_quote=strongest.get("quote_en") or strongest.get("quote_original", ""),
                        strongest_source=strongest.get("source_url", "")))
    return out


def _bare(n):
    return re.sub(r"\s*\(.*$", "", n or "").strip().lower()


def geo_index(geo):
    """Index the geography register: id lookup, children by bare name, and the country root unit(s)."""
    by_id = {g["geo_unit_id"]: g for g in geo}
    kids, roots = defaultdict(dict), {}
    for g in geo:
        if g.get("parent_id") and g["parent_id"] in by_id:
            kids[g["parent_id"]].setdefault(_bare(g["name"]), g["geo_unit_id"])
        elif g.get("unit_type") == "country":
            roots[_bare(g["name"])] = g["geo_unit_id"]
    return by_id, kids, roots


def resolve_location(location, idx):
    """Walk an actor's 'Country > Region > Place' path down the register.
    Returns (unit_id or '', number of segments matched, the segments left over as plain place text)."""
    by_id, kids, roots = idx
    parts = [p.strip() for p in (location or "").split(">") if p.strip()]
    uid = roots.get(_bare(parts[0])) if parts else None
    if not uid:
        return "", 0, parts
    depth = 1
    for seg in parts[1:]:
        nxt = kids.get(uid, {}).get(_bare(seg))
        if not nxt:
            break
        uid, depth = nxt, depth + 1
    return uid, depth, parts[depth:]


def type_list(r):
    """Primary type first, then any extra types (other_types is semicolon-separated)."""
    extra = [x.strip() for x in (r.get("other_types") or "").split(";") if x.strip()]
    return [r["actor_type"]] + [x for x in extra if x != r["actor_type"]]


def coverage(db, derived):
    grid = defaultdict(Counter)
    for r in derived:
        if r["support_status"] == "supporter":
            for ty in type_list(r):
                grid[ty][r["level"]] += 1
    return grid


def sweep_progress(db):
    """Count leaf units (areas nothing else sits under) by sweep status, grouped by their region."""
    geo = db["geo_units"]
    by_id = {g["geo_unit_id"]: g for g in geo}
    parents = {g["parent_id"] for g in geo if g.get("parent_id")}
    prog = defaultdict(Counter)
    for g in geo:
        if g["level"] in ("national", "supranational") or g["geo_unit_id"] in parents:
            continue
        p, seen, group = g, set(), "(top level)"
        while p.get("parent_id") and p["parent_id"] in by_id and p["parent_id"] not in seen:
            seen.add(p["parent_id"])
            p = by_id[p["parent_id"]]
            if p["level"] in ("region", "devolved"):
                group = p["name"]
                break
        prog[group][g["sweep_status"]] += 1
    return prog


def cmd_status(a):
    need_dir(a.dir)
    db = load_all(a.dir)
    d = derive(db)
    sup = [r for r in d if r["support_status"] == "supporter"]
    print(f"Actors: {len(sup)} supporters, {len(d) - len(sup)} former supporters | "
          f"evidence rows: {len(db['evidence'])} | links: {len(db['links'])} | exclusions: {len(db['exclusions'])}")
    g = Counter(r["best_grade"] for r in sup)
    print(f"Grades: A explicit {g['A']}, B named {g['B']}, C appears to support {g['C']}")
    print("Kinds: " + ", ".join(f"{k} {v}" for k, v in Counter(r['kind'] for r in sup).most_common()))
    levels = ["supranational", "national", "regional", "local"]
    grid = coverage(db, d)
    types = sorted(ENUMS[("actors", "actor_type")])
    print("\nCoverage grid (supporters by type × level):")
    print(f"  {'actor_type':24}" + "".join(f"{l[:9]:>10}" for l in levels) + f"{'total':>8}")
    for t in types:
        row = grid.get(t, Counter())
        print(f"  {t:24}" + "".join(f"{row[l]:>10}" for l in levels) + f"{sum(row.values()):>8}")
    prog = sweep_progress(db)
    if prog:
        print("\nLocal sweep progress (units by status, grouped by region):")
        for reg, c in sorted(prog.items()):
            tot = sum(c.values())
            print(f"  {reg:32} done {c['done']:>4}/{tot:<4} pending {c['pending']:>4}  in_progress {c['in_progress']:>3}"
                  f"  skipped {c['skipped']:>3}")
    else:
        print("\nLocal sweep: no geography units loaded yet.")
    hq = Counter(h["status"] for h in db["harvest"])
    print(f"\nHarvest lists: queued {hq['queued']}, in_progress {hq['in_progress']}, done {hq['done']}, "
          f"failed {hq['failed']}")
    unver = sum(1 for e in db["evidence"] if e.get("verified") in ("", "n"))
    print(f"Unverified evidence rows: {unver}")
    if db["runs"]:
        last = db["runs"][-1]
        print(f"\nLast run {last['run_id']} ({last['date']}): {last['focus']} → next: {last.get('notes_next','')}")
    print("\nSuggested next targets:")
    sugg = []
    q = [h for h in db["harvest"] if h["status"] in ("queued", "in_progress")]
    if q:
        sugg.append(f"process harvest lists: " + "; ".join(f"{h['harvest_id']} {h['title']}" for h in q[:5]))
    pend = [(reg, c) for reg, c in prog.items() if c["pending"] or c["in_progress"]]
    if pend:
        pend.sort(key=lambda x: (-(x[1]["in_progress"]), -x[1]["done"]))
        sugg.append(f"local sweep: continue {pend[0][0]} ({pend[0][1]['pending']} pending)")
    elif not db["geo_units"]:
        sugg.append("load the geography register (section 6.2)")
    thin = [t for t in types if sum(grid.get(t, Counter()).values()) < 3]
    if thin:
        sugg.append("thin actor types: " + ", ".join(thin))
    if g["C"]:
        sugg.append(f"upgrade {g['C']} C-graded actors")
    if unver:
        sugg.append(f"spot-check unverified evidence ({unver} rows)")
    for s in sugg or ["all queues empty — run an update sweep for new evidence since the last run"]:
        print("  - " + s)


def cmd_export(a):
    need_dir(a.dir)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("export needs openpyxl: pip install openpyxl (the CSVs are the database and can be shared as they are)")
    db = load_all(a.dir)
    errs, _ = validate(db)
    if errs:
        print(f"Note: {len(errs)} validation error(s) — run validate. Exporting anyway.")
    d = derive(db)
    slug = os.path.basename(os.path.normpath(a.dir))
    out = a.out or os.path.join(a.dir, f"{slug}_wealth_tax_supporters.xlsx")
    actors = {r["actor_id"]: r for r in db["actors"]}
    gidx = geo_index(db["geo_units"])
    geo_count = Counter()  # supporters in each unit, counting everyone located in it or below it
    for r in d:
        if r["support_status"] != "supporter":
            continue
        uid = resolve_location(r.get("location", ""), gidx)[0]
        seen = set()
        while uid and uid in gidx[0] and uid not in seen:
            seen.add(uid); geo_count[uid] += 1; uid = gidx[0][uid].get("parent_id", "")
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title, header, rows, widths=None, url_cols=()):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(header, start=1):
            w = (widths or {}).get(h)
            if w is None:
                vals = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r)]
                w = min(max(vals) + 2, 60)
            ws.column_dimensions[get_column_letter(i)].width = w
        for h in url_cols:
            if h in header:
                ci = header.index(h) + 1
                for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    for c in row:
                        if c.value and str(c.value).startswith("http"):
                            c.hyperlink = str(c.value)
                            c.font = Font(color="1F5FBF", underline="single")
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    # Overview
    ws = wb.active
    ws.title = "Overview"
    sup = [r for r in d if r["support_status"] == "supporter"]
    g = Counter(r["best_grade"] for r in sup)
    lines = [
        ("Wealth tax supporter map", ""),
        ("Area folder", slug),
        ("Exported", dt.date.today().isoformat()),
        ("Supporters", len(sup)),
        ("Former supporters (retracted)", len(d) - len(sup)),
        ("A — explicit", g["A"]), ("B — named", g["B"]), ("C — appears to support", g["C"]),
        ("Evidence rows", len(db["evidence"])),
        ("Geography units done / total",
         f"{sum(1 for x in db['geo_units'] if x['sweep_status'] == 'done')} / {len(db['geo_units'])}"),
        ("Harvest lists done / total",
         f"{sum(1 for x in db['harvest'] if x['status'] == 'done')} / {len(db['harvest'])}"),
        ("Runs", len(db["runs"])),
        ("", ""),
        ("Definition", "Actual wealth taxation since 2010: annual net wealth tax, one-off wealth levy, billionaire "
                       "minimum tax as % of wealth, taxing unrealised gains of the very wealthy, exit taxes, broad "
                       "partial wealth taxes, defending an existing wealth tax, or an unspecified wealth tax. "
                       "Excludes mansion/property/council tax, land value tax, CGT, inheritance, income, windfall, "
                       "corporate taxes and generic 'tax the rich'."),
        ("Grade A — explicit", "Specific design (rate, threshold, bill) or a formal act: vote, sponsorship, signed "
                               "motion or letter, policy adoption, publication, testimony."),
        ("Grade B — named", "Supports 'a wealth tax' by name without specifics."),
        ("Grade C — appears to support", "Ambiguous language that probably means wealth taxation; flagged."),
        ("No ranking", "Filter the Supporters tab by grade, type, level, prominence and dates."),
    ]
    for k, v in lines:
        ws.append([k, v])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=14):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    grid = coverage(db, d)
    levels = ["supranational", "national", "regional", "local"]
    ws.append([])
    ws.append(["Coverage: supporters by type × level"] + levels + ["total"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for t in sorted(ENUMS[("actors", "actor_type")]):
        row = grid.get(t, Counter())
        ws.append([t] + [row[l] for l in levels] + [sum(row.values())])

    cols = ["actor_id", "name", "kind", "actor_type", "other_types", "role", "affiliation", "party", "country", "level", "location",
            "prominence", "best_grade", "support_label", "designs", "evidence_types", "evidence_count", "first_date",
            "latest_date", "strongest_quote", "strongest_source", "wikidata_qid", "website", "status", "notes"]
    geo_name = {x["geo_unit_id"]: x["name"] for x in db["geo_units"]}
    order = sorted(sup, key=lambda r: (-GRADE_RANK.get(r["best_grade"], 0), r["actor_type"], r["name"]))
    sheet("Supporters", cols,
          [[r.get(c, "") for c in cols] for r in order],
          widths={"strongest_quote": 70, "notes": 50, "strongest_source": 40}, url_cols=("strongest_source", "website"))
    ecols = ["evidence_id", "actor_id", "actor_name"] + SCHEMA["evidence"][2:]
    ev = sorted(db["evidence"], key=lambda e: e.get("date", ""), reverse=True)
    sheet("Evidence", ecols,
          [[actors.get(e["actor_id"], {}).get("name", "") if c == "actor_name" else e.get(c, "") for c in ecols]
           for e in ev],
          widths={"quote_original": 70, "quote_en": 70, "grading_note": 50, "source_url": 40},
          url_cols=("source_url", "archive_url"))
    gcols = SCHEMA["geo_units"][:5] + ["parent_name", "supporters"] + SCHEMA["geo_units"][6:]
    sheet("Local coverage", gcols,
          [[geo_name.get(x["parent_id"], "") if c == "parent_name" else
            geo_count.get(x["geo_unit_id"], 0) if c == "supporters" else x.get(c, "") for c in gcols]
           for x in db["geo_units"]], url_cols=("website",))
    lcols = ["link_id", "from_actor_id", "from_name", "link_type", "to_actor_id", "to_name", "evidence_id",
             "source_url", "note"]
    sheet("Links", lcols,
          [[actors.get(l["from_actor_id"], {}).get("name", "") if c == "from_name" else
            actors.get(l["to_actor_id"], {}).get("name", "") if c == "to_name" else l.get(c, "") for c in lcols]
           for l in db["links"]], url_cols=("source_url",))
    pcols = ["profile_id", "actor_id", "name", "platform", "url", "source", "checked"]
    sheet("Profiles", pcols,
          [[actors.get(p["actor_id"], {}).get("name", "") if c == "name" else p.get(c, "") for c in pcols]
           for p in db["profiles"]], url_cols=("url",))
    sheet("Harvest lists", SCHEMA["harvest"], [[h.get(c, "") for c in SCHEMA["harvest"]] for h in db["harvest"]],
          url_cols=("url",))
    sheet("Run log", SCHEMA["runs"], [[r.get(c, "") for c in SCHEMA["runs"]] for r in db["runs"]])
    former = [r for r in d if r["support_status"] == "retracted"]
    if former:  # omit the sheet while nobody has retracted
        sheet("Former supporters", cols, [[r.get(c, "") for c in cols] for r in former],
              url_cols=("strongest_source",))
    wb.save(out)
    print(f"Wrote {out}")


def cmd_merge(a):
    """Merge duplicate actor DROP into KEEP: move evidence, links and profiles, fill blanks, delete DROP."""
    need_dir(a.dir)
    db = load_all(a.dir)
    by = {r["actor_id"]: r for r in db["actors"]}
    if a.keep not in by or a.drop not in by or a.keep == a.drop:
        sys.exit("keep and drop must be two different existing actor IDs")
    keep, drop = by[a.keep], by[a.drop]
    for col in SCHEMA["actors"]:
        if col in ("actor_id", "notes", "updated") or keep.get(col):
            continue
        if drop.get(col):
            keep[col] = drop[col]
    extra = [x for x in (drop.get("other_types") or "").split(";") if x and x != keep["actor_type"]]
    if drop["actor_type"] != keep["actor_type"]:
        extra.append(drop["actor_type"])
    mine = [x for x in (keep.get("other_types") or "").split(";") if x]
    keep["other_types"] = ";".join(dict.fromkeys(mine + extra))
    if norm_name(drop["name"]) != norm_name(keep["name"]):
        keep["notes"] = (keep.get("notes", "") + f" Also listed as '{drop['name']}'.").strip()
    n = 0
    for e in db["evidence"]:
        if e["actor_id"] == a.drop:
            e["actor_id"] = a.keep; n += 1
    for l in db["links"]:
        for c in ("from_actor_id", "to_actor_id"):
            if l.get(c) == a.drop:
                l[c] = a.keep
    for p in db["profiles"]:
        if p["actor_id"] == a.drop:
            p["actor_id"] = a.keep
    for r in db["actors"]:
        if r.get("affiliation_actor_id") == a.drop:
            r["affiliation_actor_id"] = a.keep
    db["actors"] = [r for r in db["actors"] if r["actor_id"] != a.drop]
    for t_ in ("actors", "evidence", "links", "profiles"):
        save(a.dir, t_, db[t_])
    print(f"Merged {a.drop} ({drop['name']}) into {a.keep} ({keep['name']}); moved {n} evidence rows")


def _plain(s):
    import html as _html
    s = _html.unescape(re.sub(r"<[^>]+>", " ", s))
    s = unicodedata.normalize("NFKC", s).lower()
    for a_, b_ in (("‘", "'"), ("’", "'"), ("“", '"'), ("”", '"'), ("–", "-"), ("—", "-")):
        s = s.replace(a_, b_)
    return re.sub(r"\s+", " ", s).strip()


def _pdf_to_text(data):
    """Text of a PDF via poppler's pdftotext when installed, else None."""
    import shutil, subprocess, tempfile
    if not shutil.which("pdftotext"):
        return None
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(data)
        path = f.name
    try:
        return subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True).stdout
    finally:
        os.unlink(path)


def _browser_text(url):
    """Fallback for pages a plain client cannot read (403, JavaScript): skill/scripts/fetch_page.py, if playwright is installed."""
    import importlib.util, subprocess
    helper = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fetch_page.py")
    if not os.path.exists(helper) or importlib.util.find_spec("playwright") is None:
        return None
    try:
        r = subprocess.run([sys.executable, helper, url, "--text", "--no-extract"], capture_output=True, text=True, timeout=180)
    except Exception:
        return None
    return r.stdout if r.returncode == 0 and r.stdout.strip() else None


def _fetch_for_check(u, timeout):
    """Page or PDF text for a quote check: plain HTTP first, then pdftotext for PDFs, then the headless browser."""
    import urllib.request
    raw = None
    try:
        req = urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (wtdb quote check)"})
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read()
            ctype = resp.headers.get("Content-Type", "").lower()
        if raw[:4] == b"%PDF" or "pdf" in ctype:
            text = _pdf_to_text(raw)
            return _plain(text) if text else None
        return _plain(raw.decode("utf-8", "ignore"))
    except Exception:
        text = _browser_text(u)
        return _plain(text) if text else None


def _dehyphenate(page):
    """Line-end hyphens in PDFs split words ("miljardärs- skatt"); join them for a second comparison."""
    return re.sub(r"(?<=\w)- (?=[a-zåäö])", "", page)


def cmd_checkquotes(a):
    """Fetch each source page (HTML, PDF via pdftotext, walled pages via the headless browser if installed) and check that the quote appears on it."""
    import random
    need_dir(a.dir)
    rows = [e for e in load(a.dir, "evidence") if e.get("source_url")]
    if a.run:
        rows = [e for e in rows if e.get("run_id") == a.run]
    if a.ids:
        rows = [e for e in rows if e["evidence_id"] in a.ids.split(",")]
    if a.sample and len(rows) > a.sample:
        random.seed(a.seed)
        rows = random.sample(rows, a.sample)
    cache, browser_cache, ok, miss, skip = {}, {}, 0, [], []
    for e in rows:
        u = e["source_url"]
        if u not in cache:
            cache[u] = _fetch_for_check(u, a.timeout)
        page = cache[u]
        if page is None:
            skip.append(e["evidence_id"]); continue
        parts = [p for p in re.split(r"…|\.\.\.", _plain(e["quote_original"])) if len(p.split()) >= 3]

        def found(pg):
            return bool(parts) and all(any(p.strip(" .") in v for v in (pg, _dehyphenate(pg))) for p in parts)

        if not found(page):
            # some sites serve a consent or shortened variant to plain clients: try the headless browser before calling it a miss
            if u not in browser_cache:
                bt = _browser_text(u)
                browser_cache[u] = _plain(bt) if bt else None
            page = browser_cache[u] or page
        if found(page):
            ok += 1
        else:
            miss.append(e["evidence_id"])
    print(f"Checked {len(rows)} rows: {ok} found on the page, {len(miss)} NOT found, {len(skip)} skipped (unreadable: no pdftotext, blocked or error)")
    if miss:
        print("Not found:", ", ".join(miss))
    if skip:
        print("Skipped:", ", ".join(skip[:40]) + (" ..." if len(skip) > 40 else ""))
    if miss and a.strict:
        sys.exit(1)


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    s = p.add_subparsers(dest="cmd", required=True)
    x = s.add_parser("init"); x.add_argument("dir"); x.add_argument("--area", required=True); x.set_defaults(f=cmd_init)
    x = s.add_parser("add"); x.add_argument("dir"); x.add_argument("table"); x.add_argument("rows"); x.set_defaults(f=cmd_add)
    x = s.add_parser("set"); x.add_argument("dir"); x.add_argument("table"); x.add_argument("id")
    x.add_argument("pairs", nargs="+"); x.set_defaults(f=cmd_set)
    for name, fn in (("validate", cmd_validate), ("dedupe", cmd_dedupe), ("status", cmd_status)):
        x = s.add_parser(name); x.add_argument("dir"); x.set_defaults(f=fn)
    x = s.add_parser("merge"); x.add_argument("dir"); x.add_argument("keep"); x.add_argument("drop"); x.set_defaults(f=cmd_merge)
    x = s.add_parser("checkquotes"); x.add_argument("dir"); x.add_argument("--run"); x.add_argument("--ids")
    x.add_argument("--sample", type=int, default=0); x.add_argument("--seed", type=int, default=1)
    x.add_argument("--timeout", type=int, default=25); x.add_argument("--strict", action="store_true")
    x.set_defaults(f=cmd_checkquotes)
    x = s.add_parser("export"); x.add_argument("dir"); x.add_argument("--out"); x.set_defaults(f=cmd_export)
    a = p.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
